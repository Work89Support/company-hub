"""Read the supplied scorecard; generate idempotent definition-only SQL."""
import hashlib
import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
source = root.parent / 'KPI SOP' / 'KPI Programmer 2026.xlsx'
sheet = openpyxl.load_workbook(source, data_only=False)['KPI Programmer']
fields = [
 'Task ID, Scope, Acceptance Criteria, Technical analysis, Dependency/Risk, เวลาเริ่มพัฒนา',
 'Task ID, Due Date ที่ตกลง, Actual Finish, Blocker พร้อมเวลาเริ่ม-จบและหลักฐานแจ้งผู้เกี่ยวข้อง',
 'Task ID, Developer Test Result, Code Review, Test Note/Test Data, เวลา Handoff',
 'Bug ID, เวลาที่ข้อมูล Incident ครบ, SLA ที่ตกลง, Root Cause, Fix/Retest และเวลาสำเร็จ',
 'Task/Release ID, Commit/PR, Release Note, Config/Migration, Smoke Test, Actual Finish',
 'Task ID, Requester, Acceptance Criteria, Handoff, Rework และสาเหตุ, Change Request',
 'Defect/Question ID, QA Owner, SLA ที่ตกลง, เวลาได้รับ-ตอบกลับ-แก้ไข, ผล Retest',
 'Release ID, DevOps Owner, Config/Migration/Dependency, Deploy Plan, Rollback, Smoke Test',
 'Incident ID, เวลาได้รับข้อมูลครบ, Impact, Root Cause, Owner, SLA, Fix/Workaround',
 'Task ID, Metric/Spec, Data Source, Date Range/Timezone/Filter, Expected/Actual Result, Business Sign-off',
]
quote = lambda s: "'" + str(s).replace("'", "''") + "'"
rows = []
assert abs(sum(sheet.cell(r, 5).value for r in range(7, 17)) - 1) < 1e-9
for r in range(7, 17):
    name, target, weight = sheet.cell(r, 3).value, sheet.cell(r, 6).value, sheet.cell(r, 5).value
    assert sheet.cell(r, 7).value is None, 'Do not import individual actuals as definitions'
    key = f'kpi-sop-2026:PROG:{r-6:02}'
    formula = sheet.cell(r, 10).value
    meta = dict(unit='ratio', achievement_cap=1.2, source_key=key, source_file=source.name,
                source_sheet=sheet.title, source_range=f'A{r}:J{r}', source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                scorecard_year=2026, department_code='PROG', category=sheet.cell(r, 2).value,
                sop_reference=sheet.cell(r, 4).value, assessment_rule=sheet['A21'].value,
                mapping=dict(required_fields=fields[r-7], suggested_source='งานจริง / กิจกรรม / ปัญหา พร้อมหลักฐานตาม SOP',
                             gap='ระบุรายการทั้งหมดและรายการที่ผ่านเกณฑ์ในรอบเดียวกัน; หัวหน้าตรวจหลักฐานก่อนรับรอง ไม่ใช้จำนวนแท็กแทนผลจริง'))
    rows.append('('+','.join([quote(key),"'PROG'",quote(name),str(target),str(weight),quote(formula),quote(f'{source.name} · {sheet.title}!A{r}:J{r}'),quote(json.dumps(meta, ensure_ascii=False))+'::jsonb'])+')')
sql = """-- Definitions from the supplied workbook. No employee actuals are manufactured.
begin;
insert into public.kpi_definitions(legacy_key,department_code,name,target,weight,formula,source,prototype_payload)
values
""" + ',\n'.join(rows) + """
on conflict(legacy_key) where legacy_key is not null do update set
 name=excluded.name,target=excluded.target,weight=excluded.weight,formula=excluded.formula,
 source=excluded.source,prototype_payload=excluded.prototype_payload,active=true;
commit;
"""
(root/'supabase/migrations/202609090032_programmer_kpi.sql').write_text(sql)
print('Prepared 10 KPI definitions; weight 100%; no actuals imported')
